"""
gerar_docx.py — Converte TCC.md em documento .docx formatado segundo a ABNT.

A conversão é feita com `python-docx`, e não por um conversor genérico, porque
a formatação exigida pela norma — fonte, entrelinhamento, margens, recuo de
primeira linha, tabelas sem fios verticais — não é o padrão de nenhum
conversor. Gerar o documento por script mantém o texto em Markdown como fonte
única e torna a formatação reproduzível a cada alteração.

Regras aplicadas (ABNT NBR 14724 e NBR 6023):

  - Fonte Arial 12 no corpo; 10 em citações longas, tabelas e legendas
  - Entrelinhamento 1,5 no corpo; simples em citações, tabelas e legendas
  - Margens: 3 cm à esquerda e superior; 2 cm à direita e inferior
  - Recuo de 1,25 cm na primeira linha de cada parágrafo do corpo
  - Corpo justificado; títulos alinhados à esquerda
  - Tabelas no padrão de fios horizontais (superior, sob o cabeçalho, inferior)
  - Figuras centralizadas, com legenda acima e fonte abaixo
  - Referências alinhadas à esquerda, sem recuo, separadas por linha em branco

Uso:
    python docs/gerar_docx.py
    python docs/gerar_docx.py --entrada TCC.md --saida TCC.docx
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent

# Fontes das ilustrações. Tabelas e figuras derivadas dos dados processados
# citam as bases de origem; as que descrevem o sistema são do próprio autor.
FONTE_AUTOR = "Fonte: elaborado pelo autor."
FONTE_DADOS = ("Fonte: elaborado pelo autor a partir de dados do SIH/SUS (DATASUS, "
               "competências de 2023) e do Censo Demográfico 2022 (IBGE).")
FONTE_MEDICAO = "Fonte: medições do autor, geradas pelos scripts do repositório."

# Legenda e fonte de cada tabela, na ordem em que aparecem no texto. Sem esta
# lista a legenda herdaria o título da seção corrente, que raramente descreve
# o conteúdo da tabela.
LEGENDAS_TABELAS = [
    ("Fontes de dados utilizadas", FONTE_AUTOR),
    ("Endpoints implementados pela API", FONTE_AUTOR),
    ("Funil de processamento das doze competências de 2023", FONTE_DADOS),
    ("Perfil geral da base processada", FONTE_DADOS),
    ("Dez capítulos CID-10 mais frequentes", FONTE_DADOS),
    ("Distribuição das internações por faixa etária e sexo", FONTE_DADOS),
    ("Municípios ordenados pelo número absoluto de internações", FONTE_DADOS),
    ("Municípios ordenados pela taxa por 100 mil habitantes", FONTE_DADOS),
    ("Métricas territoriais do recorte, obtidas por consultas espaciais", FONTE_DADOS),
    ("Classificação de autocorrelação espacial local das taxas", FONTE_DADOS),
    ("Verificação dos requisitos não funcionais", FONTE_MEDICAO),
    ("Distribuição da latência da API sob carga sustentada", FONTE_MEDICAO),
]

# Fonte de cada figura, na ordem em que aparecem no texto
FONTES_FIGURAS = [
    FONTE_AUTOR,   # 1 arquitetura
    FONTE_AUTOR,   # 2 pipeline ETL
    FONTE_AUTOR,   # 3 sequência
    FONTE_AUTOR,   # 4 entidade-relacionamento
    FONTE_AUTOR,   # 5 estados dos filtros
    FONTE_AUTOR,   # 6 componentes do frontend
    FONTE_AUTOR,   # 7 jornada do usuário
    FONTE_DADOS,   # 8 mapa coroplético
    FONTE_DADOS,   # 9 tooltip
]

FONTE = "Arial"
TAM_CORPO = Pt(12)
TAM_MENOR = Pt(10)
RECUO_PRIMEIRA_LINHA = Cm(1.25)
LARGURA_UTIL_CM = 16.0  # 21 cm de papel A4 menos as margens de 3 e 2 cm


# ---------------------------------------------------------------- utilidades

def _definir_fonte_base(documento: Document) -> None:
    """Aplica Arial 12 ao estilo Normal, incluindo a fonte de leste-asiático."""
    estilo = documento.styles["Normal"]
    estilo.font.name = FONTE
    estilo.font.size = TAM_CORPO
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), FONTE)


def _definir_margens(documento: Document) -> None:
    for secao in documento.sections:
        secao.top_margin = Cm(3)
        secao.left_margin = Cm(3)
        secao.bottom_margin = Cm(2)
        secao.right_margin = Cm(2)


def _borda(celula_ou_tabela, posicoes: dict) -> None:
    """Define fios de borda no XML — python-docx não expõe bordas diretamente."""
    tc_pr = celula_ou_tabela._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for lado, tamanho in posicoes.items():
        elemento = borders.find(qn(f"w:{lado}"))
        if elemento is None:
            elemento = OxmlElement(f"w:{lado}")
            borders.append(elemento)
        elemento.set(qn("w:val"), "single" if tamanho else "nil")
        elemento.set(qn("w:sz"), str(tamanho or 0))
        elemento.set(qn("w:color"), "000000")


# ------------------------------------------------------- formatação de texto

# Trechos entre asteriscos duplos (negrito), asteriscos simples (itálico) e
# crases (código). A ordem importa: negrito antes de itálico.
PADRAO_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[.+?\]\(.+?\))")


def _escrever_inline(paragrafo, texto: str, tamanho=TAM_CORPO) -> None:
    """Escreve texto aplicando negrito, itálico e código do Markdown."""
    for pedaco in PADRAO_INLINE.split(texto):
        if not pedaco:
            continue

        if pedaco.startswith("**") and pedaco.endswith("**"):
            run = paragrafo.add_run(pedaco[2:-2])
            run.bold = True
        elif pedaco.startswith("*") and pedaco.endswith("*"):
            run = paragrafo.add_run(pedaco[1:-1])
            run.italic = True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            run = paragrafo.add_run(pedaco[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif pedaco.startswith("["):
            # Link Markdown: preserva o rótulo e descarta a URL, que em texto
            # impresso não é clicável e polui a leitura
            rotulo = re.match(r"\[(.+?)\]\((.+?)\)", pedaco)
            run = paragrafo.add_run(rotulo.group(1) if rotulo else pedaco)
        else:
            run = paragrafo.add_run(pedaco)

        run.font.name = FONTE if run.font.name is None else run.font.name
        if run.font.size is None:
            run.font.size = tamanho


def _paragrafo_corpo(documento: Document, texto: str, recuo=True):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if recuo:
        p.paragraph_format.first_line_indent = RECUO_PRIMEIRA_LINHA
    _escrever_inline(p, texto)
    return p


def _titulo(documento: Document, texto: str, nivel: int):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(18 if nivel <= 2 else 12)
    p.paragraph_format.space_after = Pt(12 if nivel <= 2 else 6)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(texto.upper() if nivel <= 2 else texto)
    run.font.name = FONTE
    run.font.size = TAM_CORPO
    run.bold = nivel <= 3
    run.italic = nivel >= 4
    return p


def _legenda(documento: Document, texto: str, alinhamento=WD_ALIGN_PARAGRAPH.CENTER):
    p = documento.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texto)
    run.font.name = FONTE
    run.font.size = TAM_MENOR
    return p


# ------------------------------------------------------------------ blocos

def _adicionar_tabela(documento: Document, linhas: list, contador: dict) -> None:
    """Insere uma tabela Markdown no padrão de fios horizontais da ABNT."""
    def celulas(linha):
        return [c.strip() for c in linha.strip().strip("|").split("|")]

    cabecalho = celulas(linhas[0])
    corpo = [celulas(linha) for linha in linhas[2:]]
    n_colunas = len(cabecalho)

    contador["tabela"] += 1
    indice_legenda = contador["tabela"] - 1
    if indice_legenda < len(LEGENDAS_TABELAS):
        titulo, fonte = LEGENDAS_TABELAS[indice_legenda]
    else:
        titulo, fonte = contador["titulo_tabela"], FONTE_AUTOR
    _legenda(documento, f"Tabela {contador['tabela']} — {titulo}",
             WD_ALIGN_PARAGRAPH.LEFT)

    tabela = documento.add_table(rows=1 + len(corpo), cols=n_colunas)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = True

    for indice, texto in enumerate(cabecalho):
        celula = tabela.cell(0, indice)
        celula.text = ""
        p = celula.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        _escrever_inline(p, f"**{texto}**" if texto else "", TAM_MENOR)
        # Fio superior da tabela e fio sob o cabeçalho
        _borda(celula, {"top": 8, "bottom": 8, "left": 0, "right": 0})

    for i, linha in enumerate(corpo, start=1):
        for j in range(n_colunas):
            celula = tabela.cell(i, j)
            celula.text = ""
            p = celula.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            _escrever_inline(p, linha[j] if j < len(linha) else "", TAM_MENOR)
            # Fio inferior apenas na última linha; nenhum fio vertical
            ultimo = 8 if i == len(corpo) else 0
            _borda(celula, {"top": 0, "bottom": ultimo, "left": 0, "right": 0})

    _legenda(documento, fonte, WD_ALIGN_PARAGRAPH.LEFT)


def _adicionar_figura(documento: Document, caminho: Path, descricao: str,
                      contador: dict) -> None:
    if not caminho.exists():
        print(f"[docx] figura ausente, ignorada: {caminho}", file=sys.stderr)
        return

    contador["figura"] += 1
    _legenda(documento, f"Figura {contador['figura']} — {descricao}",
             WD_ALIGN_PARAGRAPH.LEFT)

    paragrafo = documento.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo.paragraph_format.space_after = Pt(0)
    paragrafo.add_run().add_picture(str(caminho), width=Cm(LARGURA_UTIL_CM))

    indice = contador["figura"] - 1
    fonte = FONTES_FIGURAS[indice] if indice < len(FONTES_FIGURAS) else FONTE_AUTOR
    _legenda(documento, fonte, WD_ALIGN_PARAGRAPH.LEFT)


def _adicionar_referencia(documento: Document, texto: str) -> None:
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    _escrever_inline(p, texto)


def _adicionar_lista(documento: Document, texto: str) -> None:
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = RECUO_PRIMEIRA_LINHA
    p.paragraph_format.space_after = Pt(0)
    _escrever_inline(p, f"• {texto}")


def _adicionar_citacao(documento: Document, texto: str) -> None:
    """Nota ou citação longa: recuo de 4 cm, corpo 10, entrelinhamento simples."""
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _escrever_inline(p, texto, TAM_MENOR)


# ------------------------------------------------------------------ conversão

PADRAO_IMAGEM = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
PADRAO_TITULO = re.compile(r"^(#{1,6})\s+(.*)$")


def converter(entrada: Path, saida: Path) -> None:
    linhas = entrada.read_text(encoding="utf-8").split("\n")

    documento = Document()
    _definir_fonte_base(documento)
    _definir_margens(documento)

    contador = {"figura": 0, "tabela": 0, "titulo_tabela": "dados do estudo"}
    em_referencias = False
    i = 0

    while i < len(linhas):
        linha = linhas[i]
        despida = linha.strip()

        # Linhas em branco e separadores horizontais não geram parágrafo
        if not despida or despida == "---":
            i += 1
            continue

        # Blocos de código: preservados em monoespaçada, sem justificação
        if despida.startswith("```"):
            i += 1
            bloco = []
            while i < len(linhas) and not linhas[i].strip().startswith("```"):
                bloco.append(linhas[i])
                i += 1
            i += 1
            p = documento.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run("\n".join(bloco))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            continue

        titulo = PADRAO_TITULO.match(despida)
        if titulo:
            nivel = len(titulo.group(1))
            texto = titulo.group(2)
            em_referencias = texto.strip().lower().startswith("referências")
            # O título do trabalho (nível 1) recebe centralização e destaque
            if nivel == 1:
                p = documento.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(24)
                run = p.add_run(texto.upper())
                run.font.name = FONTE
                run.font.size = Pt(14)
                run.bold = True
            else:
                _titulo(documento, texto, nivel)
                # O texto que antecede uma tabela costuma nomeá-la; guarda-se o
                # título corrente para compor a legenda
                contador["titulo_tabela"] = texto
            i += 1
            continue

        imagem = PADRAO_IMAGEM.match(despida)
        if imagem:
            descricao, caminho = imagem.group(1), imagem.group(2)
            _adicionar_figura(documento, RAIZ / caminho, descricao, contador)
            i += 1
            continue

        # Tabela: linha de cabeçalho seguida da linha de separação
        if despida.startswith("|") and i + 1 < len(linhas) and \
                set(linhas[i + 1].strip()) <= set("|-: "):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            _adicionar_tabela(documento, bloco, contador)
            continue

        if despida.startswith(">"):
            _adicionar_citacao(documento, despida.lstrip("> ").strip())
            i += 1
            continue

        if despida.startswith(("- ", "* ")):
            _adicionar_lista(documento, despida[2:])
            i += 1
            continue

        item_numerado = re.match(r"^(\d+)\.\s+(.*)$", despida)
        if item_numerado:
            p = documento.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = RECUO_PRIMEIRA_LINHA
            p.paragraph_format.space_after = Pt(0)
            _escrever_inline(p, f"{item_numerado.group(1)}. {item_numerado.group(2)}")
            i += 1
            continue

        if em_referencias:
            _adicionar_referencia(documento, despida)
        else:
            # O nome do autor, logo abaixo do título, é centralizado
            recuo = not despida.startswith("**André")
            p = _paragrafo_corpo(documento, despida, recuo=recuo)
            if not recuo:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1

    documento.save(saida)
    print(f"[docx] gerado: {saida}")
    print(f"[docx] {contador['figura']} figuras e {contador['tabela']} tabelas")


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera o .docx do TCC a partir do Markdown.")
    parser.add_argument("--entrada", default=str(RAIZ / "TCC.md"))
    parser.add_argument("--saida", default=str(RAIZ / "TCC.docx"))
    args = parser.parse_args()

    converter(Path(args.entrada), Path(args.saida))


if __name__ == "__main__":
    main()
